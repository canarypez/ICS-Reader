"""用户上传的课程资料（PDF / PPTX）。

和 ICS 课件不同，这些文件没有 lecturekit 渲染出来的 `lecture.json`，所以这里负责：

* 把用户选的文件复制到 `DATA_DIR/documents/<12 位 id>/source.<ext>`；
* 逐页抽出文字（PDF 用 pypdf；PPTX 先用 PowerPoint/WPS 转成 PDF，转不了再用
  python-pptx 做一份简化的 HTML 幻灯片）；
* 维护每页的标题 / 页码，让 `agent.build_context` 能像 ICS 课件一样喂给模型；
* 按需把 PDF 页渲染成 PNG（pypdfium2 / PDFium），供 `web/docview.html` 显示。

存放布局：

    documents/
      <id>/
        source.pdf | source.pptx   原件副本
        converted.pdf              PPTX 转换出的 PDF（有 Office/WPS 时）
        doc.json                   元数据 + 每页文本
        slides.json                无 Office 时的简化 HTML 幻灯片
        assets/                    简化渲染用到的图片
        cache/                     页面 PNG 缓存

注意：`doc.json` 里存的是**用户自己上传的资料的文字**，只在本机；程序不会上传它，
只有提问时抽出的上下文会发给用户配置的模型接口。
"""
import html
import json
import os
import re
import shutil
import subprocess
import threading
import time
import uuid
from io import BytesIO

import pypdfium2 as pdfium
from PIL import Image
from pypdf import PdfReader
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

import agent
import config as cfgmod
from sync import CREATE_NO_WINDOW

#: 文档 id 只允许小写十六进制，既够随机，也让所有路径参数都能直接校验。
DOC_ID_RE = re.compile(r"^[0-9a-f]{12}$")

#: 简化渲染时支持的图片格式。EMF/WMF 之类浏览器不认，直接跳过。
_IMG_EXT = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"}

#: PDFium 不是线程安全的。多个请求同时要不同页时串行渲染，反正单页通常几十毫秒。
_PDF_LOCK = threading.Lock()


def _remove_tree(path: str) -> bool:
    """删除目录并确认结果；Windows 上文件被占用时不要假装删成功了。"""
    if not path or not os.path.exists(path):
        return True

    def _retry(func, target, _exc):
        try:
            os.chmod(target, 0o700)
        except OSError:
            pass
        try:
            func(target)
        except OSError:
            pass

    for _ in range(4):
        try:
            shutil.rmtree(path, onexc=_retry)
        except OSError:
            pass
        if not os.path.exists(path):
            return True
        time.sleep(0.12)
    return not os.path.exists(path)

#: 侧栏里一次最多展示多少页的目录。太长的 PDF 全给出来只会把上下文撑爆；
#: 当前页和前后页始终会单独给出，检索也能找到没列出来的页。
_OUTLINE_LIMIT = 120


def documents_root() -> str:
    return os.path.join(cfgmod.DATA_DIR, "documents")


# --------------------------------------------------------------------------- 文本
def _clean_text(text) -> str:
    """统一换行、压掉行内多余空白和连续空行。"""
    text = str(text or "").replace("\r\n", "\n").replace("\r", "\n")
    out = []
    blank = False
    for line in text.split("\n"):
        line = re.sub(r"[ \t]+", " ", line).strip()
        if not line:
            if out and not blank:
                out.append("")
            blank = True
        else:
            out.append(line)
            blank = False
    return "\n".join(out).strip()


def _first_line(text, limit=60) -> str:
    for line in (text or "").split("\n"):
        line = line.strip()
        if line:
            return line[:limit]
    return ""


def _page_title(text, number) -> str:
    return _first_line(text) or f"第 {number} 页"


def _doc_tree(pages) -> list:
    tree = []
    for p in pages[:_OUTLINE_LIMIT]:
        label = f"第 {p.get('number') or '?'} 页"
        title = (p.get("title") or "").strip()
        if title:
            label += f" · {title}"
        tree.append({"title": label})
    if len(pages) > _OUTLINE_LIMIT:
        tree.append({"title": f"……（其余 {len(pages) - _OUTLINE_LIMIT} 页略）"})
    return tree


# --------------------------------------------------------------------------- PDF
def _extract_pdf(path: str, log=lambda s: None) -> list:
    """用 pypdf 抽每页文字。扫描件抽不出来是正常的，页面仍然可以看。"""
    try:
        reader = PdfReader(path)
    except Exception as e:
        raise RuntimeError(f"读不了这份 PDF：{e}")
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            pass
        if reader.is_encrypted:
            raise RuntimeError("这份 PDF 有密码保护，请先去掉密码再上传。")

    total = len(reader.pages)
    pages = []
    for i in range(total):
        try:
            text = _clean_text(reader.pages[i].extract_text() or "")
        except Exception:
            text = ""
        number = i + 1
        pages.append({
            "id": f"p{number}",
            "number": number,
            "title": _page_title(text, number)[:80],
            "text": text,
        })
        if total > 40 and (i + 1) % 20 == 0:
            log(f"已读取 {i + 1}/{total} 页…")
    return pages


def _pdf_outline(path: str) -> list:
    """PDF 书签 → [{title,index,depth}]。没有书签就返回空，调用方退化成每页首行。"""
    try:
        reader = PdfReader(path)
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                return []
        raw = reader.outline
    except Exception:
        return []
    out = []

    def walk(items, depth=0):
        for item in items or []:
            if isinstance(item, list):
                walk(item, depth + 1)
                continue
            try:
                title = _clean_text(getattr(item, "title", "") or "")
                page = reader.get_destination_page_number(item)
            except Exception:
                continue
            if title and isinstance(page, int) and page >= 0:
                out.append({"title": title[:120], "index": page, "depth": depth})

    try:
        walk(raw)
    except Exception:
        return []
    return out[:400]


# --------------------------------------------------------------------------- Markdown
def _read_text_file(path: str) -> str:
    with open(path, "rb") as f:
        raw = f.read()
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def _md_title(text: str, fallback: str) -> str:
    for line in text.split("\n"):
        m = re.match(r"^\s*#\s+(.+?)\s*#*\s*$", line)
        if m:
            return m.group(1).strip()[:120]
    return fallback


def _extract_md(path: str):
    """Markdown 按二级标题分页；没有二级标题就按 `---`，再没有就整篇一页。"""
    text = _read_text_file(path).replace("\r\n", "\n").replace("\r", "\n")
    # 去掉 YAML front matter（--- 开头、--- 结尾），它不该出现在正文里
    front = re.match(r"^\s*---\n([\s\S]*?)\n---\s*(?:\n|$)", text)
    if front:
        text = text[front.end():]
    title = _md_title(text, os.path.splitext(os.path.basename(path))[0])

    starts = [m.start() for m in re.finditer(r"^##\s+\S", text, re.M)]
    if not starts:
        starts = [m.start() for m in re.finditer(r"^---\s*$", text, re.M)]
    chunks = []
    if not starts:
        chunks = [text]
    else:
        if starts[0] > 0:
            chunks.append(text[:starts[0]])
        for i, s in enumerate(starts):
            e = starts[i + 1] if i + 1 < len(starts) else len(text)
            chunks.append(text[s:e])

    pages, outline = [], []
    n = 0
    for chunk in chunks:
        chunk = chunk.strip("\n")
        if not chunk.strip():
            continue
        n += 1
        head = ""
        for line in chunk.split("\n"):
            m = re.match(r"^\s*#{1,6}\s+(.+)", line)
            if m:
                head = m.group(1).strip()
                break
        if not head:
            head = _first_line(chunk, 80)
        pages.append({
            "id": f"p{n}", "number": n, "title": head[:80],
            "text": chunk, "notes": "",
        })
        outline.append({"title": head[:120], "index": n - 1, "depth": 0})
    return pages, (outline or _doc_tree(pages))


# --------------------------------------------------------------------------- PPTX
def _convert_pptx(src: str, dst: str) -> tuple:
    """调用本机 PowerPoint / WPS 把 PPTX 转成 PDF。

    返回 (成功?, 说明)。没有 Office 时返回 False，调用方退回简化 HTML 渲染。
    用 PowerShell 的 COM 而不是 pywin32：程序打包后不必再带一个 COM 绑定库。
    """
    script = cfgmod.resource("pptx2pdf.ps1")
    if not os.path.isfile(script):
        return False, "找不到 pptx2pdf.ps1，跳过 PowerPoint 转换"
    exe = shutil.which("powershell")
    if not exe:
        exe = os.path.join(os.environ.get("SystemRoot", r"C:\Windows"),
                           "System32", "WindowsPowerShell", "v1.0", "powershell.exe")
    if not os.path.isfile(exe):
        return False, "找不到 PowerShell，跳过 PowerPoint 转换"

    try:
        p = subprocess.Popen(
            [exe, "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
             "-File", script, "-Src", src, "-Dst", dst],
            stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True,
            encoding="utf-8", errors="replace", creationflags=CREATE_NO_WINDOW,
        )
        try:
            out, _ = p.communicate(timeout=180)
        except subprocess.TimeoutExpired:
            subprocess.run(["taskkill", "/PID", str(p.pid), "/T", "/F"],
                           capture_output=True, creationflags=CREATE_NO_WINDOW)
            return False, "PowerPoint/WPS 转换超时"
    except OSError as e:
        return False, f"无法启动 PowerShell：{e}"

    if p.returncode == 0 and os.path.isfile(dst) and os.path.getsize(dst) > 0:
        return True, ""
    detail = (out or "").strip().splitlines()
    return False, (detail[-1] if detail else f"退出码 {p.returncode}")[:400]


def _iter_shapes(shapes):
    """递归展开组合形状，按 z 序产出。"""
    for sh in shapes:
        try:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from _iter_shapes(sh.shapes)
                continue
        except Exception:
            pass
        yield sh


def _shape_text(sh) -> str:
    parts = []
    try:
        if getattr(sh, "has_text_frame", False):
            text = sh.text_frame.text
            if text:
                parts.append(text)
    except Exception:
        pass
    try:
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
    except Exception:
        pass
    try:
        if getattr(sh, "has_chart", False) and sh.has_chart:
            chart = sh.chart
            if chart.has_title and chart.chart_title.has_text_frame:
                parts.append(chart.chart_title.text_frame.text)
            categories = [str(c) for c in chart.plots[0].categories]
            if categories:
                parts.append(" / ".join(categories))
            for series in chart.series:
                name = str(series.name or "")
                values = [str(v) for v in series.values]
                parts.append(f"{name}: " + ", ".join(values))
    except Exception:
        pass
    return "\n".join(p for p in parts if p)


def _extract_slide_title(slide) -> str:
    try:
        title = slide.shapes.title
        if title is not None:
            return (title.text or "").strip()
    except Exception:
        pass
    return ""


# ---- 简化 HTML 渲染（没有 PowerPoint / WPS 时的退路） ---------------------
def _solid_fill(sh):
    try:
        fill = sh.fill
        if fill.type == 1:                      # MSO_FILL.SOLID
            rgb = fill.fore_color.rgb
            if rgb is not None:
                return "#" + str(rgb)
    except Exception:
        pass
    return None


def _bg_color(slide) -> str:
    for obj in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        try:
            fill = obj.background.fill
            if fill.type == 1:
                rgb = fill.fore_color.rgb
                if rgb is not None:
                    return "#" + str(rgb)
        except Exception:
            pass
    return "#ffffff"


def _text_color(bg: str) -> str:
    """深色底上用白字。主题色取不到 RGB 时也能猜个大概。"""
    try:
        h = bg.lstrip("#")
        r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
        if (0.299 * r + 0.587 * g + 0.114 * b) / 255 < 0.45:
            return "#ffffff"
    except Exception:
        pass
    return "#222222"


def _is_title(sh) -> bool:
    try:
        if sh.is_placeholder:
            return sh.placeholder_format.type in (
                PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
    except Exception:
        pass
    return False


def _default_pt(sh) -> float:
    try:
        if sh.is_placeholder:
            t = sh.placeholder_format.type
            if t in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                return 40
            if t == PP_PLACEHOLDER.SUBTITLE:
                return 24
    except Exception:
        pass
    return 18


def _tf_html(tf, default_pt, text_color, is_title, px_per_pt, px_per_emu) -> str:
    paras = []
    for p in tf.paragraphs:
        text = p.text or ""
        if not text and not p.runs:
            continue
        align = "left"
        try:
            if p.alignment == PP_ALIGN.CENTER:
                align = "center"
            elif p.alignment == PP_ALIGN.RIGHT:
                align = "right"
            elif p.alignment == PP_ALIGN.JUSTIFY:
                align = "justify"
        except Exception:
            pass

        level = p.level or 0
        bullet = ""
        if text.strip() and not is_title:
            bullet = "• " if level == 0 else ("– " if level == 1 else "· ")

        runs = []
        for r in p.runs:
            t = r.text or ""
            if not t:
                continue
            styles = []
            try:
                f = r.font
                if f.size is not None:
                    styles.append(f"font-size:{f.size.pt * px_per_pt:.1f}px")
                if f.bold:
                    styles.append("font-weight:700")
                if f.italic:
                    styles.append("font-style:italic")
                if f.underline:
                    styles.append("text-decoration:underline")
                try:
                    rgb = f.color.rgb
                    if rgb is not None:
                        styles.append(f"color:#{rgb}")
                except Exception:
                    pass
            except Exception:
                pass
            runs.append(f'<span style="{";".join(styles)}">{html.escape(t)}</span>')
        if not runs:
            if not text:
                continue
            runs = [html.escape(text)]

        pstyle = [f"text-align:{align}", f"margin-left:{level * 20}px"]
        try:
            ls = p.line_spacing
            if isinstance(ls, float):
                pstyle.append(f"line-height:{ls:.2f}")
            elif ls is not None:
                pstyle.append(f"line-height:{ls.pt * px_per_pt:.1f}px")
        except Exception:
            pass
        try:
            if p.space_before is not None:
                pstyle.append(f"margin-top:{p.space_before.pt * px_per_pt:.1f}px")
            if p.space_after is not None:
                pstyle.append(f"margin-bottom:{p.space_after.pt * px_per_pt:.1f}px")
        except Exception:
            pass
        paras.append(f'<div style="{";".join(pstyle)}">{bullet}{"".join(runs)}</div>')

    if not paras:
        return ""

    anchor = "flex-start"
    try:
        if tf.vertical_anchor == MSO_ANCHOR.MIDDLE:
            anchor = "center"
        elif tf.vertical_anchor == MSO_ANCHOR.BOTTOM:
            anchor = "flex-end"
    except Exception:
        pass

    padding = []
    for name, value in (("left", tf.margin_left), ("top", tf.margin_top),
                        ("right", tf.margin_right), ("bottom", tf.margin_bottom)):
        try:
            padding.append(f"padding-{name}:{value * px_per_emu:.1f}px")
        except Exception:
            pass

    base = [
        "display:flex", "flex-direction:column", f"justify-content:{anchor}",
        "height:100%", "box-sizing:border-box",
        f"color:{text_color}", f"font-size:{default_pt * px_per_pt:.1f}px",
        "font-family:'Microsoft YaHei','PingFang SC','Segoe UI',sans-serif",
        "line-height:1.25", "white-space:pre-wrap", "word-break:break-word",
    ]
    return f'<div style="{";".join(base + padding)}">{"".join(paras)}</div>'


def _chart_html(chart) -> str:
    parts = []
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            parts.append(
                '<div style="font-weight:700;margin-bottom:4px">'
                + html.escape(chart.chart_title.text_frame.text) + "</div>")
    except Exception:
        pass
    try:
        categories = [str(c) for c in chart.plots[0].categories]
        head = "<tr><th></th>" + "".join(
            f"<th>{html.escape(c)}</th>" for c in categories) + "</tr>"
        rows = [head]
        for series in chart.series:
            name = html.escape(str(series.name or ""))
            values = [str(v) for v in series.values]
            rows.append(f"<tr><th>{name}</th>" + "".join(
                f"<td>{html.escape(v)}</td>" for v in values) + "</tr>")
        parts.append(
            '<table style="border-collapse:collapse;width:100%;font-size:12px">'
            + "".join(rows) + "</table>")
    except Exception:
        parts.append('<div style="opacity:.6">（图表）</div>')
    return "".join(parts)


def _shape_box(sh):
    try:
        return (sh.left or 0, sh.top or 0, sh.width or 0, sh.height or 0)
    except Exception:
        return (0, 0, 0, 0)


def _group_xf(group):
    """组合形状的子坐标 → 父坐标。读 group 的 xfrm（off/ext/chOff/chExt）。"""
    try:
        xfrm = group._element.grpSpPr.xfrm
        off, ext = xfrm.off, xfrm.ext
        ch_off, ch_ext = xfrm.chOff, xfrm.chExt
        sx = float(ext.cx or 0) / float(ch_ext.cx or 1)
        sy = float(ext.cy or 0) / float(ch_ext.cy or 1)
        return lambda box: (
            off.x + (box[0] - ch_off.x) * sx,
            off.y + (box[1] - ch_off.y) * sy,
            box[2] * sx,
            box[3] * sy,
        )
    except Exception:
        return None


def _compose(outer, inner):
    if outer is None:
        return inner
    if inner is None:
        return outer
    return lambda box: outer(inner(box))


def _render_shape(sh, box, ctx) -> str:
    left, top, width, height = box
    scale = ctx["scale"]
    x, y = left * scale, top * scale
    w, h = width * scale, height * scale
    if w <= 0 and h <= 0:
        return ""
    pos = (f"position:absolute;left:{x:.1f}px;top:{y:.1f}px;"
           f"width:{w:.1f}px;height:{h:.1f}px;")

    # 图片（普通图片和装了图片的占位符都走这里）
    try:
        image = sh.image
        ext = (image.ext or "").lower()
        if ext in _IMG_EXT and image.blob:
            name = f"img{ctx['idx']}.{ext}"
            ctx["idx"] += 1
            with open(os.path.join(ctx["assets"], name), "wb") as f:
                f.write(image.blob)
            return (f'<img alt="" src="/api/doc/{ctx["doc_id"]}/asset/{name}" '
                    f'style="{pos}object-fit:contain">')
    except Exception:
        pass

    # 表格
    try:
        if getattr(sh, "has_table", False) and sh.has_table:
            rows = []
            for row in sh.table.rows:
                cells = []
                for cell in row.cells:
                    cells.append(
                        '<td style="border:1px solid #c9ced6;padding:4px 6px;'
                        'vertical-align:top">' + html.escape(cell.text) + "</td>")
                rows.append("<tr>" + "".join(cells) + "</tr>")
            return (f'<div style="{pos}overflow:hidden"><table style="border-collapse:'
                    f'collapse;width:100%;font-size:{ctx["px_per_pt"] * 14:.1f}px;'
                    f'font-family:\'Microsoft YaHei\',sans-serif">'
                    + "".join(rows) + "</table></div>")
    except Exception:
        pass

    # 图表：退化成数据表，至少让学生看到数
    try:
        if getattr(sh, "has_chart", False) and sh.has_chart:
            return (f'<div style="{pos}overflow:hidden;font-family:\'Microsoft YaHei\','
                    f'sans-serif">{_chart_html(sh.chart)}</div>')
    except Exception:
        pass

    fill = _solid_fill(sh)

    # 文本框
    try:
        if getattr(sh, "has_text_frame", False):
            inner = _tf_html(
                sh.text_frame, _default_pt(sh), ctx["text_color"],
                _is_title(sh), ctx["px_per_pt"], ctx["px_per_emu"])
            if inner:
                bg = f";background:{fill}" if fill else ""
                return f'<div style="{pos}overflow:hidden{bg}">{inner}</div>'
    except Exception:
        pass

    # 只有填充色的装饰块 / 线条
    if fill and (w > 0 or h > 0):
        w2, h2 = max(w, 2), max(h, 2)
        return (f'<div style="position:absolute;left:{x:.1f}px;top:{y:.1f}px;'
                f'width:{w2:.1f}px;height:{h2:.1f}px;background:{fill}"></div>')
    try:
        if sh.shape_type == MSO_SHAPE_TYPE.LINE:
            return (f'<div style="position:absolute;left:{x:.1f}px;'
                    f'top:{y + h / 2:.1f}px;width:{max(w, 2):.1f}px;height:2px;'
                    f'background:#888888"></div>')
    except Exception:
        pass
    return ""


def _render_shapes(shapes, xf, ctx, skip_placeholder=False):
    for sh in shapes:
        try:
            if skip_placeholder:
                try:
                    if sh.is_placeholder:
                        continue
                except Exception:
                    pass
            box = _shape_box(sh)
            if xf:
                box = xf(box)
            try:
                if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                    _render_shapes(sh.shapes, _compose(xf, _group_xf(sh)), ctx, False)
                    continue
            except Exception:
                pass
            piece = _render_shape(sh, box, ctx)
            if piece:
                ctx["out"].append(piece)
        except Exception:
            continue


def _render_slide(slide, slide_w, slide_h, doc_id, assets, out_w=1280) -> str:
    scale = out_w / float(slide_w or 1)
    out_h = max(1, int(round(out_w * float(slide_h or 6858000)
                             / float(slide_w or 12192000))))
    bg = _bg_color(slide)
    ctx = {
        "out": [], "idx": 0, "doc_id": doc_id, "assets": assets,
        "scale": scale, "px_per_emu": scale, "px_per_pt": 12700.0 * scale,
        "text_color": _text_color(bg),
    }
    # 先画母版/版式里的装饰（跳过占位符，免得把「单击此处编辑」之类画上来），
    # 再画本页内容。这样背景图、校徽这类东西不会丢。
    for owner, skip in ((slide.slide_layout.slide_master, True),
                        (slide.slide_layout, True),
                        (slide, False)):
        try:
            _render_shapes(owner.shapes, None, ctx, skip_placeholder=skip)
        except Exception:
            continue
    # 固定 1280×out_h 的画布：前端只需整体等比缩放，不必知道原始 slide 尺寸。
    return (f'<div style="position:relative;width:{out_w}px;height:{out_h}px;'
            f'background:{bg};overflow:hidden">'
            + "".join(ctx["out"]) + "</div>")


def _slide_notes(slide) -> str:
    """Speaker Notes。读不到（旧格式、损坏）就当没有，不能让它阻断导入。"""
    try:
        frame = slide.notes_slide.notes_text_frame
        if frame is None:
            return ""
        return _clean_text(frame.text or "")
    except Exception:
        return ""


def _merge_pptx_notes(pdf_pages, pptx_pages):
    """PPTX 转成 PDF 后，文字以 PDF 为准，但备注只有 PPTX 里有。"""
    by_num = {p.get("number"): p for p in (pptx_pages or [])}
    for p in pdf_pages:
        other = by_num.get(p.get("number"))
        if not other:
            continue
        p["notes"] = other.get("notes") or ""
        if other.get("title") and not p.get("title"):
            p["title"] = other["title"]
    return pdf_pages


def _extract_pptx(src: str, out_dir: str, doc_id: str, log=lambda s: None,
                  render_slides=True):
    try:
        prs = Presentation(src)
    except Exception as e:
        raise RuntimeError(f"读不了这份 PPTX：{e}")
    assets = os.path.join(out_dir, "assets")
    os.makedirs(assets, exist_ok=True)
    slide_w = prs.slide_width or 12192000
    slide_h = prs.slide_height or 6858000

    pages, slides = [], []
    for i, slide in enumerate(prs.slides):
        number = i + 1
        parts = []
        for sh in _iter_shapes(slide.shapes):
            text = _shape_text(sh)
            if text:
                parts.append(text)
        text = _clean_text("\n".join(parts))
        title = _extract_slide_title(slide) or _page_title(text, number)
        pages.append({
            "id": f"p{number}",
            "number": number,
            "title": title[:80],
            "text": text,
            "notes": _slide_notes(slide),
        })
        if render_slides:
            slides.append({
                "title": title[:80],
                "html": _render_slide(slide, slide_w, slide_h, doc_id, assets),
            })
        if number % 10 == 0:
            log(f"已处理 {number} 张幻灯片…")
    return pages, slides


# --------------------------------------------------------------------------- 存取
class DocStore:
    """用户资料库。接口尽量对齐 agent.CourseStore 的 `lectures` 形状。

    可以给多个 root：多课程时每个 `courses/<cid>/items/` 是一个 root，迁移失败的
    旧 `documents/` 再挂在最后。同一个 id 以先扫到的为准（课程目录优先）。
    """

    def __init__(self, root, write_root: str = None):
        # 兼容老调用：root 传字符串时当成唯一 root
        self.roots = [root] if isinstance(root, str) else list(root or [])
        self.write_root = write_root or (self.roots[0] if self.roots else "")
        self.bases = {}          # doc_id -> 实际目录
        self.lectures = {}
        self.reload()

    # ---- 读取 ----
    def reload(self):
        self.lectures = {}
        self.bases = {}
        for root in self.roots:
            if not os.path.isdir(root):
                continue
            for name in sorted(os.listdir(root)):
                if not DOC_ID_RE.match(name):
                    continue
                if name in self.lectures:
                    continue
                base = os.path.join(root, name)
                entry = self._load(base, name)
                if entry:
                    self.lectures[name] = entry
                    self.bases[name] = base
        return self.lectures

    def _load(self, base: str, doc_id: str):
        try:
            with open(os.path.join(base, "doc.json"), "r", encoding="utf-8") as f:
                data = json.load(f)
        except Exception:
            return None
        pages = [p for p in (data.get("pages") or [])
                 if isinstance(p, dict) and p.get("id")]
        if not pages:
            return None
        slides = []
        if data.get("mode") == "slide":
            try:
                with open(os.path.join(base, "slides.json"), "r", encoding="utf-8") as f:
                    slides = json.load(f) or []
            except Exception:
                slides = []
        page_text, page_number, page_title = {}, {}, {}
        for i, p in enumerate(pages):
            pid = p["id"]
            page_text[pid] = p.get("text") or ""
            page_number[pid] = p.get("number") or (i + 1)
            page_title[pid] = p.get("title") or ""
        outline = data.get("outline") or []
        if not outline:
            outline = _doc_tree(pages)
        return {
            "dir": doc_id,
            "vid": doc_id,
            "id": doc_id,
            "kind": "doc",
            "course_id": data.get("course_id") or "",
            "doc_type": data.get("type") or "pdf",
            "mode": data.get("mode") or "pdf",
            "title": data.get("title") or data.get("name") or doc_id,
            "name": data.get("name") or "",
            "added": data.get("added") or 0,
            "tree": outline,
            "outline": outline,
            "pages": pages,
            "page_text": page_text,
            "page_number": page_number,
            "page_title": page_title,
            "slides": slides,
            "has_text": any(t.strip() for t in page_text.values()),
            "base": f"/docview.html?doc={doc_id}",
        }

    def get(self, doc_id):
        return self.lectures.get(doc_id)

    def payload_one(self, doc_id):
        d = self.lectures.get(doc_id)
        if not d:
            return None
        return {
            "id": doc_id, "vid": doc_id, "title": d["title"], "name": d["name"],
            "kind": "doc", "type": d["doc_type"], "mode": d["mode"],
            "course_id": d.get("course_id") or "",
            "page_count": len(d["pages"]), "has_text": d["has_text"],
            "base": d.get("base") or f"/docview.html?doc={doc_id}",
        }

    def payload(self):
        items = [self.payload_one(k) for k in self.lectures]
        items = [x for x in items if x]
        items.sort(key=lambda x: self.lectures[x["id"]].get("added") or 0, reverse=True)
        return items

    # ---- 导入 ----
    def add(self, src: str, name: str = None, log=lambda s: None,
            course_id: str = "", out_root: str = None):
        src = os.path.abspath(str(src or ""))
        if not src or not os.path.isfile(src):
            raise RuntimeError("找不到这个文件。")
        original = os.path.basename(name or src) or "upload"
        ext = os.path.splitext(original)[1].lower()
        if ext not in (".pdf", ".pptx", ".md"):
            raise RuntimeError("只支持 PDF、PPTX 和 Markdown 文件。")

        doc_id = uuid.uuid4().hex[:12]
        root = out_root or self.write_root
        if not root:
            raise RuntimeError("没有可写入的资料目录")
        base = os.path.join(root, doc_id)
        os.makedirs(os.path.join(base, "assets"), exist_ok=True)
        os.makedirs(os.path.join(base, "cache"), exist_ok=True)
        try:
            stored = os.path.join(base, "source" + ext)
            shutil.copyfile(src, stored)

            mode = "pdf"
            view = stored
            slides = []
            outline = []
            if ext == ".pdf":
                log("正在读取 PDF 文字…")
                pages = _extract_pdf(stored, log)
                outline = _pdf_outline(stored)
            elif ext == ".md":
                log("正在解析 Markdown…")
                pages, outline = _extract_md(stored)
                mode = "md"
                view = stored
            else:
                pdf = os.path.join(base, "converted.pdf")
                log("正在把 PPTX 转成 PDF（调用本机 PowerPoint/WPS）…")
                ok, detail = _convert_pptx(stored, pdf)
                log("正在读取幻灯片文字与备注…")
                pptx_pages, slides = _extract_pptx(
                    stored, base, doc_id, log, render_slides=not ok)
                if ok:
                    log("转换完成，正在读取文字…")
                    pdf_pages = _extract_pdf(pdf, log)
                    pages = _merge_pptx_notes(pdf_pages, pptx_pages)
                    view = pdf
                    slides = []
                else:
                    log("未能调用 PowerPoint/WPS，改用内置简化渲染。")
                    if detail:
                        log(detail)
                    mode = "slide"
                    pages = pptx_pages
                    with open(os.path.join(base, "slides.json"), "w",
                              encoding="utf-8") as f:
                        json.dump(slides, f, ensure_ascii=False)
            if not pages:
                raise RuntimeError("这份文件里没有可读取的页面。")
            if mode == "slide" and len(slides) != len(pages):
                raise RuntimeError("幻灯片渲染不完整，请换一份文件重试。")
            if not outline:
                outline = _doc_tree(pages)

            data = {
                "id": doc_id,
                "name": original[:200],
                "title": (os.path.splitext(original)[0] or "未命名资料")[:120],
                "type": ext.lstrip("."),
                "mode": mode,
                "course_id": course_id or os.path.basename(os.path.dirname(root)),
                "added": time.time(),
                "pages": pages,
                "outline": outline,
            }
            with open(os.path.join(base, "doc.json"), "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=1)
            entry = self._load(base, doc_id)
            if entry:
                self.lectures[doc_id] = entry
                self.bases[doc_id] = base
                if root not in self.roots:
                    self.roots.append(root)
            return self.payload_one(doc_id)
        except Exception:
            shutil.rmtree(base, ignore_errors=True)
            raise

    def delete(self, doc_id: str) -> bool:
        if not DOC_ID_RE.match(str(doc_id or "")):
            return False
        # 同一份旧资料可能同时存在于 courses/<cid>/items/（迁移副本）和
        # documents/（迁移原件）。只删当前 base 的话，refresh 时另一个 root
        # 会把它重新扫出来，看起来就是「删了又回来」。所有 root 里的同 id
        # 目录都要清掉。
        targets = []
        base = self.bases.get(doc_id) or os.path.join(self.write_root, doc_id)
        if os.path.isdir(base):
            targets.append(base)
        for root in self.roots:
            cand = os.path.join(root, doc_id)
            if os.path.isfile(os.path.join(cand, "doc.json")):
                targets.append(cand)
        if not targets:
            return False
        ok = True
        for path in dict.fromkeys(os.path.abspath(p) for p in targets):
            if not _remove_tree(path):
                ok = False
        self.lectures.pop(doc_id, None)
        self.bases.pop(doc_id, None)
        return ok

    # ---- 页面 ----
    def _pdf_path(self, doc_id):
        base = self.bases.get(doc_id) or os.path.join(self.write_root, doc_id)
        converted = os.path.join(base, "converted.pdf")
        if os.path.isfile(converted):
            return converted
        source = os.path.join(base, "source.pdf")
        return source if os.path.isfile(source) else None

    def render_page(self, doc_id: str, number: int, width: int = 1600) -> bytes:
        entry = self.get(doc_id)
        if not entry:
            raise RuntimeError("资料不存在")
        if entry["mode"] != "pdf":
            raise RuntimeError("这份资料没有 PDF 页面")
        number = int(number)
        if number < 1 or number > len(entry["pages"]):
            raise RuntimeError("页码超出范围")
        pdf_path = self._pdf_path(doc_id)
        if not pdf_path:
            raise RuntimeError("原始 PDF 不见了")

        width = max(200, min(int(width or 1600), 2800))
        width = int(round(width / 200.0) * 200)
        cache_dir = os.path.join(self.bases.get(doc_id) or self.write_root, "cache")
        os.makedirs(cache_dir, exist_ok=True)
        cache = os.path.join(cache_dir, f"p{number}-{width}.png")
        if os.path.isfile(cache):
            with open(cache, "rb") as f:
                return f.read()

        with _PDF_LOCK:
            if os.path.isfile(cache):
                with open(cache, "rb") as f:
                    return f.read()
            pdf = pdfium.PdfDocument(pdf_path)
            try:
                page = pdf[number - 1]
                page_w = page.get_size()[0] or 1
                scale = max(0.4, min(width / float(page_w), 4.0))
                bitmap = page.render(scale=scale)
                image = bitmap.to_pil()
                if image.mode in ("RGBA", "LA"):
                    flat = Image.new("RGB", image.size, (255, 255, 255))
                    flat.paste(image, mask=image.split()[-1])
                    image = flat
                buf = BytesIO()
                image.save(buf, format="PNG")
                data = buf.getvalue()
            finally:
                try:
                    pdf.close()
                except Exception:
                    pass
        tmp = cache + ".tmp"
        with open(tmp, "wb") as f:
            f.write(data)
        os.replace(tmp, cache)
        return data

    # ---- 透明文本层 ----
    _MONO_FONT = re.compile(r"(mono|courier|consolas|menlo|monaco|code|source\s*code)",
                            re.I)
    _MATH_FONT = re.compile(r"(math|cmu|cambria\s*math|stix|symbol)", re.I)
    _MATH_CHARS = set("∑∏∫√∞≈≠≤≥±×÷∂∇∈∉⊂⊃∪∩→←↔⇒⇔∀∃αβγδεζηθικλμνξπρστυφχψωΓΔΘΛΞΠΣΦΨΩ")

    def text_layer(self, doc_id: str, number: int):
        """PDF 页的字符框，前端在 PNG 上盖一层透明可选文字。

        PDFium 的原点是左下角，这里统一换成 top-left；坐标单位是 PDF 点（72dpi），
        前端按百分比定位，所以图片渲染宽度怎么变都不需要重算。
        """
        entry = self.get(doc_id)
        if not entry or entry["mode"] != "pdf":
            return None
        number = int(number)
        if number < 1 or number > len(entry["pages"]):
            return None
        pdf_path = self._pdf_path(doc_id)
        if not pdf_path:
            return None
        base = self.bases.get(doc_id) or os.path.join(self.write_root, doc_id)
        cache_dir = os.path.join(base, "textlayer")
        cache = os.path.join(cache_dir, f"p{number}.json")
        if os.path.isfile(cache):
            try:
                with open(cache, "r", encoding="utf-8") as f:
                    return json.load(f)
            except Exception:
                pass

        with _PDF_LOCK:
            pdf = pdfium.PdfDocument(pdf_path)
            try:
                page = pdf[number - 1]
                page_w, page_h = page.get_size()
                tp = page.get_textpage()
                try:
                    data = self._layer_from_textpage(tp, page_w, page_h, number)
                finally:
                    try:
                        tp.close()
                    except Exception:
                        pass
            finally:
                try:
                    pdf.close()
                except Exception:
                    pass

        os.makedirs(cache_dir, exist_ok=True)
        tmp = cache + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False)
        os.replace(tmp, cache)
        return data

    def _layer_from_textpage(self, tp, page_w, page_h, number):
        lines = []
        cur = []

        def flush():
            if not cur:
                return
            text = "".join(c["t"] for c in cur)
            if not text.strip():
                cur.clear()
                return
            x0 = min(c["x"] for c in cur)
            y0 = min(c["y"] for c in cur)
            x1 = max(c["x"] + c["w"] for c in cur)
            y1 = max(c["y"] + c["h"] for c in cur)
            spans = []
            for c in cur:
                prev = spans[-1] if spans else None
                if prev and prev["style"] == c["style"] and \
                        c["x"] - (prev["x"] + prev["w"]) <= max(1.0, c["size"] * 0.12):
                    # 相邻且同样式：并成一个 span，少几千个 DOM 节点
                    nx = min(prev["x"], c["x"])
                    ny = min(prev["y"], c["y"])
                    nr = max(prev["x"] + prev["w"], c["x"] + c["w"])
                    nb = max(prev["y"] + prev["h"], c["y"] + c["h"])
                    prev.update(x=nx, y=ny, w=nr - nx, h=nb - ny, text=prev["text"] + c["t"])
                else:
                    spans.append({
                        "x": round(c["x"], 2), "y": round(c["y"], 2),
                        "w": round(c["w"], 2), "h": round(c["h"], 2),
                        "text": c["t"], "font": c["font"], "size": round(c["size"], 2),
                        "style": c["style"],
                    })
            # 合并后的 span 也要把坐标收圆
            for sp in spans:
                if "text" in sp and "x" in sp:
                    sp["x"] = round(sp["x"], 2)
                    sp["y"] = round(sp["y"], 2)
                    sp["w"] = round(sp["w"], 2)
                    sp["h"] = round(sp["h"], 2)
            lines.append({
                "x": round(x0, 2), "y": round(y0, 2),
                "w": round(x1 - x0, 2), "h": round(y1 - y0, 2),
                "text": text, "spans": spans,
            })
            cur.clear()

        count = tp.count_chars()
        last_cy = None
        last_h = 0.0
        for i in range(count):
            try:
                ch = tp.get_text_range(i, 1)
            except Exception:
                continue
            if not ch:
                continue
            if ch in ("\r", "\n"):
                flush()
                last_cy = None
                continue
            try:
                left, bottom, right, top = tp.get_charbox(i)
            except Exception:
                continue
            w, h = right - left, top - bottom
            if w <= 0 or h <= 0:
                # 空格有时没有 box；当成一个正常空白，保留在行内
                if ch.isspace() and cur:
                    cur.append({"t": " ", "x": cur[-1]["x"] + cur[-1]["w"],
                                "y": cur[-1]["y"], "w": max(1.0, h), "h": h,
                                "font": cur[-1]["font"], "size": cur[-1]["size"],
                                "style": cur[-1]["style"]})
                continue
            y = page_h - top
            cy = y + h / 2.0
            if cur and last_cy is not None and \
                    abs(cy - last_cy) > max(2.0, 0.62 * max(h, last_h)):
                flush()
            font, size, style = "", h, "text"
            try:
                obj = tp.get_textobj(i)
                try:
                    size = float(obj.get_font_size()) or h
                except Exception:
                    size = h
                try:
                    font = obj.get_font().get_family_name() or ""
                except Exception:
                    font = ""
            except Exception:
                pass
            if self._MONO_FONT.search(font or ""):
                style = "code"
            elif self._MATH_FONT.search(font or "") or ch in self._MATH_CHARS:
                style = "math"
            cur.append({
                "t": ch, "x": left, "y": y, "w": w, "h": h,
                "font": font, "size": size, "style": style,
            })
            last_cy = cy
            last_h = h
        flush()
        return {"number": number, "width": round(page_w, 2),
                "height": round(page_h, 2), "lines": lines}

    def page_notes(self, doc_id: str, number: int) -> str:
        entry = self.get(doc_id)
        if not entry:
            return ""
        try:
            n = int(number)
        except Exception:
            return ""
        if n < 1 or n > len(entry["pages"]):
            return ""
        return entry["pages"][n - 1].get("notes") or ""

    def outline(self, doc_id: str):
        entry = self.get(doc_id)
        return (entry or {}).get("outline") or []

    def page_text(self, doc_id: str, number: int) -> str:
        entry = self.get(doc_id)
        if not entry:
            return ""
        try:
            number = int(number)
        except Exception:
            return ""
        if number < 1 or number > len(entry["pages"]):
            return ""
        return entry["page_text"].get(entry["pages"][number - 1]["id"], "")

    def page_slide(self, doc_id: str, number: int):
        entry = self.get(doc_id)
        if not entry or entry["mode"] != "slide":
            return None
        try:
            number = int(number)
        except Exception:
            return None
        if number < 1 or number > len(entry["slides"]):
            return None
        return entry["slides"][number - 1]

    def asset_path(self, doc_id: str, name: str):
        if not DOC_ID_RE.match(str(doc_id or "")):
            return None
        name = os.path.basename(str(name or ""))
        if not re.fullmatch(r"[A-Za-z0-9_.-]+", name):
            return None
        base = self.bases.get(doc_id) or os.path.join(self.write_root, doc_id)
        path = os.path.join(base, "assets", name)
        return path if os.path.isfile(path) else None


class Library(agent.CourseStore):
    """所有课程（ICS / 仓库 / 上传资料）合并成的知识库。

    继承了 CourseStore 的检索和上下文方法 —— 它们只依赖 `self.lectures` 这个字典，
    所以把各课程按同样的形状塞进去就行了；`course_id` 用来把检索限制在一门课里。
    """

    def __init__(self, registry):
        self.registry = registry
        self.build_root = registry.build_root
        self.docs = DocStore(registry.doc_roots())
        self.lectures = {}
        self.reload()

    def reload(self):
        self.lectures = {}
        for cid in self.registry.course_ids():
            course = self.registry.load(cid) or {}
            if cid == "ics":
                # ICS 保持老路径、老 lid（目录名）和历史文件名
                root = self.build_root
                if not os.path.isdir(root):
                    continue
                for name in sorted(os.listdir(root)):
                    entry = self.load_lecture(
                        os.path.join(root, name), name,
                        base=f"/lectures/{name}/", course_id=cid)
                    if entry:
                        self.lectures[name] = entry
                continue
            root = self.registry.render_root(cid)
            if not os.path.isdir(root):
                continue
            for name in sorted(os.listdir(root)):
                lid = f"{cid}:{name}"
                entry = self.load_lecture(
                    os.path.join(root, name), lid,
                    base=f"/course/{cid}/lectures/{name}/", course_id=cid,
                    vid=f"{cid}:{name}")
                if entry:
                    self.lectures[lid] = entry

        for did, entry in self.docs.reload().items():
            if not entry.get("course_id"):
                entry["course_id"] = "mydocs"
            self.lectures[did] = entry

    def refresh_docs(self):
        """只重读资料，不动课件（上传/删除后调用）。课程目录也可能变，整树重扫一次。"""
        self.reload()

    def course_items(self, cid):
        return [self.lectures[lid] for lid in self.lecture_ids()
                if self.lectures[lid].get("course_id") == cid]
